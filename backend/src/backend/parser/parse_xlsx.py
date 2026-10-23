from datetime import date, datetime, time
from typing import BinaryIO

from openpyxl import load_workbook

from backend.parser.parsing_errors import EmptyDocument, ParseError


def _fmt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, (datetime, date, time)):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def parse_xlsx(f: BinaryIO) -> str:
    try:
        wb = load_workbook(
            f,
            read_only=True,
            data_only=True,
        )
    except Exception as e:
        raise ParseError(f"cannot open xlsx: {e}") from e

    sections = []

    try:
        for ws in wb.worksheets:
            if ws.sheet_state != "visible":
                continue

            iterator = ws.iter_rows(values_only=True)
            header = None
            rows = []

            for row in iterator:
                cells = [_fmt(c) for c in row]

                # Ignore completely empty rows.
                if not any(cells):
                    continue

                # Use the first non-empty row as the header.
                if header is None:
                    while cells and not cells[-1]:
                        cells.pop()

                    header = []
                    used_headers = set()

                    for i, name in enumerate(cells, start=1):
                        name = name or f"Column_{i}"
                        original_name = name
                        suffix = 2

                        while name in used_headers:
                            name = f"{original_name}_{suffix}"
                            suffix += 1

                        header.append(name)
                        used_headers.add(name)

                    continue

                # Preserve the column-to-value relationships.
                if len(cells) > len(header):
                    for i in range(len(header) + 1, len(cells) + 1):
                        header.append(f"Column_{i}")

                # Pad shorter rows to match the header length.
                cells += [""] * (len(header) - len(cells))

                fields = [
                    f"{name}: {value}"
                    for name, value in zip(header, cells)
                    if name and value
                ]

                if fields:
                    rows.append(" | ".join(fields))

            if rows:
                sections.append(f"## Sheet: {ws.title}\n" + "\n".join(rows))

    except Exception as e:
        raise ParseError(f"cannot read xlsx: {e}") from e
    finally:
        wb.close()

    if not sections:
        raise EmptyDocument("xlsx has no data")

    return "\n\n".join(sections)
